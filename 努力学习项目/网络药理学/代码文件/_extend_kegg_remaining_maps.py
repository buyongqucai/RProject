# -*- coding: utf-8 -*-
"""把 KEGG 上仍缺的 map（无人源 hsa 对应）追加进清单与中文名表。幂等。"""
from __future__ import annotations

import csv
import re
from pathlib import Path

import requests

MAN = Path(r"E:\RProject\努力学习项目\网络药理学\交付文件\图片\KEGG官方通路图\00_全部通路图清单.csv")
ZH_PY = Path(r"E:\RProject\努力学习项目\网络药理学\代码文件\kegg_zh_names.py")
XLSX = Path(r"D:\数据库\KEGG数据库\KEGG通路动画映射表.xlsx")

ZH = {
    "map01220": "芳香族化合物降解",
    "map01310": "氮循环",
    "map03230": "病毒基因组结构",
    "map03240": "病毒复制",
    "map03262": "病毒颗粒-冠状病毒",
    "map03261": "病毒颗粒-流感病毒",
    "map03263": "病毒颗粒-肠道病毒",
    "map03268": "病毒颗粒-正痘病毒",
    "map03259": "病毒颗粒-噬菌体T4与T7",
    "map03258": "病毒颗粒-λ噬菌体",
    "map02060": "磷酸转移酶系统(PTS)",
    "map03070": "细菌分泌系统",
    "map02020": "双组分系统",
    "map04013": "MAPK信号通路（果蝇）",
    "map04016": "MAPK信号通路（植物）",
    "map04011": "MAPK信号通路（酵母）",
    "map04341": "Hedgehog信号通路（果蝇）",
    "map04391": "Hippo信号通路（果蝇）",
    "map04075": "植物激素信号转导",
    "map04149": "过氧化物酶体（酵母）",
    "map04138": "自噬（酵母）",
    "map04139": "线粒体自噬（酵母）",
    "map04111": "细胞周期（酵母）",
    "map04112": "细胞周期（Caulobacter）",
    "map04113": "减数分裂（酵母）",
    "map04214": "凋亡（果蝇）",
    "map02024": "群体感应",
    "map05111": "生物膜形成-霍乱弧菌",
    "map02025": "生物膜形成-铜绿假单胞菌",
    "map02026": "生物膜形成-大肠杆菌",
    "map02030": "细菌趋化性",
    "map02040": "鞭毛组装",
    "map04624": "Toll与Imd信号通路",
    "map04745": "光传导（果蝇）",
    "map04320": "背腹轴形成",
    "map04361": "轴突再生",
    "map04212": "长寿调节通路（线虫）",
    "map04711": "昼夜节律（果蝇）",
    "map04712": "昼夜节律（植物）",
    "map04626": "植物-病原互作",
    "map01501": "β-内酰胺抗性",
    "map01502": "万古霉素抗性",
    "map01503": "阳离子抗菌肽(CAMP)抗性",
    "map07011": "青霉素类",
    "map07012": "头孢菌素-注射剂",
    "map07013": "头孢菌素-口服剂",
    "map07021": "氨基糖苷类",
    "map07019": "四环素类",
    "map07020": "大环内酯与酮内酯",
    "map07014": "喹诺酮类",
    "map07023": "利福霉素类",
    "map07026": "抗真菌药",
    "map07044": "抗病毒药",
    "map07053": "抗HIV药",
    "map07040": "抗肿瘤药-烷化剂",
    "map07041": "抗肿瘤药-抗代谢药",
    "map07042": "抗肿瘤药-天然产物",
    "map07043": "抗肿瘤药-激素类",
    "map07045": "抗肿瘤药-蛋白激酶抑制剂",
    "map07032": "催眠药",
    "map07030": "抗焦虑药",
    "map07033": "抗惊厥药",
    "map07015": "局部镇痛药",
    "map07039": "阿片类镇痛药",
    "map07028": "抗精神病药",
    "map07029": "抗精神病药-吩噻嗪类",
    "map07031": "抗精神病药-丁酰苯类",
    "map07027": "抗抑郁药",
    "map07056": "阿尔茨海默病治疗药",
    "map07057": "抗帕金森药",
    "map07055": "磺胺衍生物总览",
    "map07016": "磺胺衍生物-磺胺药",
    "map07017": "磺胺衍生物-利尿药",
    "map07018": "磺胺衍生物-降糖药",
    "map07037": "抗心律失常药",
    "map07038": "抗溃疡药",
    "map07046": "免疫抑制剂",
    "map07047": "抗骨质疏松药",
    "map07048": "抗偏头痛药",
    "map07049": "抗血栓药",
    "map07050": "抗风湿药-DMARD与生物制剂",
    "map07051": "抗糖尿病药",
    "map07052": "调血脂药",
    "map07054": "抗青光眼药",
    "map07220": "胆碱能与抗胆碱能药",
    "map07215": "α-肾上腺素能受体激动剂与拮抗剂",
    "map07214": "β-肾上腺素能受体激动剂与拮抗剂",
    "map07213": "多巴胺受体激动剂与拮抗剂",
    "map07212": "组胺H1受体拮抗剂",
    "map07227": "组胺H2/H3受体激动剂与拮抗剂",
    "map07211": "5-羟色胺受体激动剂与拮抗剂",
    "map07228": "类二十烷酸受体激动剂与拮抗剂",
    "map07224": "阿片受体激动剂与拮抗剂",
    "map07229": "血管紧张素与内皮素受体拮抗剂",
    "map07225": "糖皮质与盐皮质激素受体激动剂与拮抗剂",
    "map07226": "孕激素雄激素与雌激素受体激动剂与拮抗剂",
    "map07223": "维A酸受体与RXR激动剂与拮抗剂",
    "map07222": "PPAR激动剂",
    "map07221": "烟碱型胆碱受体拮抗剂",
    "map07230": "GABA-A受体激动剂与拮抗剂",
    "map07036": "钙通道阻滞药",
    "map07231": "钠通道阻滞药",
    "map07232": "钾通道阻滞与开放药",
    "map07235": "NMDA受体拮抗剂",
    "map07233": "离子转运体抑制剂",
    "map07234": "神经递质转运体抑制剂",
    "map07216": "儿茶酚胺转移酶抑制剂",
    "map07219": "环氧化酶抑制剂",
    "map07024": "HMG-CoA还原酶抑制剂",
    "map07217": "肾素-血管紧张素系统抑制剂",
    "map07218": "HIV蛋白酶抑制剂",
    "map07025": "喹啉类",
    "map07034": "类二十烷酸",
    "map07035": "前列腺素",
    "map07110": "苯甲酸类",
    "map07112": "1,2-二苯基取代类",
    "map07114": "萘类",
    "map07117": "苯二氮䓬类",
}


def fname_en(en: str) -> str:
    s = re.sub(r"[^A-Za-z0-9\- ]", "", en).replace(" ", "_")
    return re.sub(r"_+", "_", s).strip("_")


def main() -> None:
    s = requests.Session()
    s.trust_env = False
    s.headers["User-Agent"] = "Mozilla/5.0"
    have: set[str] = set()
    max_rank = 0
    with MAN.open(encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            have.add(str(r["pathway_id"]).strip())
            max_rank = max(max_rank, int(r["rank"]))
    hsa_codes = set()
    for line in s.get("https://rest.kegg.jp/list/pathway/hsa", timeout=60).text.splitlines():
        hsa_codes.add(line.split("\t", 1)[0].replace("path:", "")[3:])
    todo: list[tuple[str, str]] = []
    for line in s.get("https://rest.kegg.jp/list/pathway", timeout=60).text.splitlines():
        pid, name = line.split("\t", 1)
        pid = pid.replace("path:", "")
        if pid.startswith("map") and pid not in have and pid[3:] not in hsa_codes:
            todo.append((pid, name.strip()))
    missing_zh = [p for p, _ in todo if p not in ZH]
    if missing_zh:
        raise SystemExit(f"ZH dict missing: {missing_zh}")

    rows = []
    for pid, en in todo:
        max_rank += 1
        png = f"{max_rank:03d}_{pid}_{fname_en(en)}.png"
        url = f"https://rest.kegg.jp/get/{pid}/image"
        rows.append((max_rank, pid, ZH[pid], en, png, url))

    with MAN.open("a", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        for rank, pid, zh, en, png, url in rows:
            w.writerow([rank, pid, en, png, url])

    src = ZH_PY.read_text(encoding="utf-8")
    add = []
    for rank, pid, zh, en, png, url in rows:
        if f'"{pid}"' not in src:
            add.append(f'    "{pid}": "{zh}",')
    if add:
        src = src.rstrip()
        assert src.endswith("}")
        src = src[:-1].rstrip() + "\n" + "\n".join(add) + "\n}\n"
        ZH_PY.write_text(src, encoding="utf-8")

    x_added = 0
    if XLSX.exists():
        from openpyxl import load_workbook
        wb = load_workbook(XLSX)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        i_pid = header.index("KEGG编号")
        existing = {r[i_pid] for r in ws.iter_rows(min_row=2, values_only=True)}
        for rank, pid, zh, en, png, url in rows:
            if pid in existing:
                continue
            gif = f"{rank:03d}_{pid}_{zh}_{fname_en(en)}_动画.gif"
            ws.append([rank, pid, zh, en, "待绘制", gif, png, url])
            x_added += 1
        wb.save(XLSX)

    print(f"added {len(rows)} maps ranks {rows[0][0]}-{rows[-1][0]} xlsx+{x_added}")
    for rank, pid, zh, en, png, url in rows[:12]:
        print(f"{rank}\t{pid}\t{zh}\t{en}")


if __name__ == "__main__":
    main()
