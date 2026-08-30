# -*- coding: utf-8 -*-
"""KEGG 官方通路图批量下载：按映射表「待绘制」行抓 PNG 到 KEGG通路图。

映射表：D:\\数据库\\KEGG数据库\\KEGG通路动画映射表.xlsx
  列：排名 / KEGG编号 / 通路中文名 / 通路英文名 / 动画状态 / 动画GIF / 官方通路图PNG / KEGG官方图URL
仅下载「动画状态 == 待绘制」且本地尚无 PNG 的行；已存在跳过（幂等）。
"""
from __future__ import annotations

import time
from pathlib import Path

import requests
from openpyxl import load_workbook

XLSX = Path(r"D:\数据库\KEGG数据库\KEGG通路动画映射表.xlsx")
OUT = Path(r"D:\数据库\KEGG数据库\KEGG通路图")

wb = load_workbook(XLSX, read_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
header = rows[0]
i_status = header.index("动画状态")
i_png = header.index("官方通路图PNG")
i_url = header.index("KEGG官方图URL")

todo = [(r[i_png], r[i_url]) for r in rows[1:]
        if r[i_status] == "待绘制" and r[i_png] and r[i_url]]
print(f"待绘制 {len(todo)} 张", flush=True)

s = requests.Session()
s.headers.update({"User-Agent": "Mozilla/5.0"})
ok = skip = fail = 0
for name, url in todo:
    dest = OUT / name
    if dest.exists() and dest.stat().st_size > 1000:
        skip += 1
        continue
    try:
        r = s.get(url, timeout=60)
        r.raise_for_status()
        if not r.content.startswith(b"\x89PNG"):
            raise ValueError("not a PNG")
        dest.write_bytes(r.content)
        ok += 1
        print(f"[ok] {name} {len(r.content)//1024}KB", flush=True)
    except Exception as e:
        fail += 1
        print(f"[fail] {name} {e}", flush=True)
    time.sleep(0.4)
print(f"done: ok={ok} skip={skip} fail={fail}", flush=True)
