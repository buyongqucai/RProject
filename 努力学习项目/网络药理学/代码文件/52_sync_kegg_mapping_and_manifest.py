# -*- coding: utf-8 -*-
"""同步 KEGG 映射表与合成清单：
1) 映射表 95 行（rank<=183 且 GIF 已存在）动画状态 → 已绘制
2) 新增 184-213 行的 GIF 文件名改用无括号中文短名
3) 追加 30 行到 00_全部通路图清单.csv；追加 zh 名到 kegg_zh_names.py
幂等：重复运行不产生重复行。
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(r"D:\数据库\KEGG数据库\KEGG通路动画映射表.xlsx")
GIF_DIR = Path(r"D:\数据库\KEGG数据库\KEGG动画图")
MAN = Path(r"E:\RProject\努力学习项目\网络药理学\交付文件\图片\KEGG官方通路图\00_全部通路图清单.csv")
ZH_PY = Path(r"E:\RProject\努力学习项目\网络药理学\代码文件\kegg_zh_names.py")

ZH_NEW = {
    "hsa05200": "癌症通路总览", "hsa04216": "铁死亡", "hsa04142": "溶酶体",
    "hsa04141": "内质网蛋白质加工", "hsa04144": "胞吞作用", "hsa04146": "过氧化物酶体",
    "hsa04330": "Notch信号通路", "hsa04340": "Hedgehog信号通路",
    "hsa04512": "ECM受体相互作用", "hsa04612": "抗原加工与呈递",
    "hsa04022": "cGMP-PKG信号通路", "hsa04710": "昼夜节律", "hsa04714": "产热作用",
    "hsa04724": "谷氨酸能突触", "hsa04727": "GABA能突触", "hsa04728": "多巴胺能突触",
    "hsa04911": "胰岛素分泌", "hsa04974": "蛋白质消化与吸收", "hsa04978": "矿物质吸收",
    "hsa04979": "胆固醇代谢", "hsa05012": "帕金森病", "hsa05014": "肌萎缩侧索硬化",
    "hsa05016": "亨廷顿病", "hsa05410": "肥厚型心肌病", "hsa05414": "扩张型心肌病",
    "hsa04260": "心肌收缩", "hsa04261": "心肌细胞肾上腺素能信号",
    "hsa00100": "类固醇生物合成", "hsa04975": "脂肪消化与吸收",
    "hsa04977": "维生素消化与吸收",
}


def fen(en: str) -> str:
    s = re.sub(r"[^A-Za-z0-9\- ]", "", en).replace(" ", "_")
    return re.sub(r"_+", "_", s)


wb = load_workbook(XLSX)
ws = wb.active
header = [c.value for c in ws[1]]
i_rank = header.index("排名") + 1
i_pid = header.index("KEGG编号") + 1
i_zh = header.index("通路中文名") + 1
i_en = header.index("通路英文名") + 1
i_st = header.index("动画状态") + 1
i_gif = header.index("动画GIF") + 1
i_png = header.index("官方通路图PNG") + 1
i_url = header.index("KEGG官方图URL") + 1

existing_gifs = {p.name for p in GIF_DIR.glob("*.gif")}
n_flip = 0
new_rows = []
for row in range(2, ws.max_row + 1):
    rank = ws.cell(row, i_rank).value
    pid = ws.cell(row, i_pid).value
    en = ws.cell(row, i_en).value
    if rank is None or pid is None:
        continue
    if rank <= 183:
        gif = ws.cell(row, i_gif).value
        if ws.cell(row, i_st).value == "待绘制" and gif in existing_gifs:
            ws.cell(row, i_st).value = "已绘制"
            n_flip += 1
    else:
        zh = ZH_NEW.get(pid, ws.cell(row, i_zh).value)
        ws.cell(row, i_zh).value = zh
        gif_name = f"{rank:03d}_{pid}_{zh}_{fen(en)}_动画.gif"
        ws.cell(row, i_gif).value = gif_name
        new_rows.append((rank, pid, en, zh))
wb.save(XLSX)
print(f"flipped 已绘制: {n_flip}; new rows: {len(new_rows)}")

# manifest 追加
with MAN.open(encoding="utf-8-sig") as fh:
    man_rows = list(csv.DictReader(fh))
man_ranks = {int(r["rank"]) for r in man_rows}
with MAN.open("a", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    for rank, pid, en, zh in new_rows:
        if rank in man_ranks:
            continue
        png = f"{rank:03d}_{pid}_{fen(en)}.png"
        w.writerow([rank, pid, en, png, f"https://rest.kegg.jp/get/{pid}/image"])
print(f"manifest total: {max(man_ranks | {r[0] for r in new_rows})}")

# kegg_zh_names.py 追加
src = ZH_PY.read_text(encoding="utf-8")
add = []
for pid, zh in ZH_NEW.items():
    if f'"{pid}"' not in src:
        add.append(f'    "{pid}": "{zh}",')
if add:
    src = src.rstrip()
    assert src.endswith("}")
    src = src[:-1].rstrip() + "\n" + "\n".join(add) + "\n}\n"
    ZH_PY.write_text(src, encoding="utf-8")
print(f"zh names added: {len(add)}")
