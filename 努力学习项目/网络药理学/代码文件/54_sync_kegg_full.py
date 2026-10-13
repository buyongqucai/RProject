# -*- coding: utf-8 -*-
"""同步 53_extend_kegg_full.py 扩展后的映射表（rank>=214）：
1) 追加新行到 00_全部通路图清单.csv（45/47/48 号脚本的输入清单）；
2) 追加 pid→中文名 到 kegg_zh_names.py（48 号合成脚本取中文名用）。
幂等：已存在的 rank / pid 跳过。
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(r"D:\数据库\KEGG数据库\KEGG通路动画映射表.xlsx")
MAN = Path(r"E:\RProject\努力学习项目\网络药理学\交付文件\图片\KEGG官方通路图\00_全部通路图清单.csv")
ZH_PY = Path(r"E:\RProject\努力学习项目\网络药理学\代码文件\kegg_zh_names.py")


def main() -> None:
    wb = load_workbook(XLSX, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = list(rows[0])
    i_rank = header.index("排名")
    i_pid = header.index("KEGG编号")
    i_zh = header.index("通路中文名")
    i_en = header.index("通路英文名")
    i_png = header.index("官方通路图PNG")
    i_url = header.index("KEGG官方图URL")

    new_rows = []
    for r in rows[1:]:
        if r[i_rank] is None or r[i_pid] is None:
            continue
        new_rows.append((int(r[i_rank]), str(r[i_pid]), str(r[i_zh] or ""),
                         str(r[i_en] or ""), str(r[i_png] or ""), str(r[i_url] or "")))

    # manifest 追加
    with MAN.open(encoding="utf-8-sig") as fh:
        man_ranks = {int(r["rank"]) for r in csv.DictReader(fh)}
    added_man = 0
    with MAN.open("a", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        for rank, pid, zh, en, png, url in new_rows:
            if rank in man_ranks:
                continue
            w.writerow([rank, pid, en, png, url])
            added_man += 1
    print(f"manifest added: {added_man} (total {len(man_ranks) + added_man})")

    # kegg_zh_names.py 追加
    src = ZH_PY.read_text(encoding="utf-8")
    add = []
    for rank, pid, zh, en, png, url in new_rows:
        if zh and f'"{pid}"' not in src:
            add.append(f'    "{pid}": "{zh}",')
    if add:
        src = src.rstrip()
        assert src.endswith("}")
        src = src[:-1].rstrip() + "\n" + "\n".join(add) + "\n}\n"
        ZH_PY.write_text(src, encoding="utf-8")
    print(f"zh names added: {len(add)}")


if __name__ == "__main__":
    main()
