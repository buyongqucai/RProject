# -*- coding: utf-8 -*-
"""KEGG 映射表扩展：补常见/前沿通路（官方名校验自 rest.kegg.jp），追加排名并下载官方 PNG。

- 候选清单人工策展（常见信号/细胞过程/疾病/心代谢前沿），与现有 183 条去重。
- 官方英文名以 https://rest.kegg.jp/list/pathway/hsa 为准（校验 ID 存在）。
- 追加行：动画状态=待绘制；PNG/GIF 文件名按既有命名范式生成；URL=rest.kegg.jp/get/{id}/image。
- 幂等：已在映射表的 ID 跳过。
"""
from __future__ import annotations

import re
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

XLSX = Path(r"D:\数据库\KEGG数据库\KEGG通路动画映射表.xlsx")
OUT = Path(r"D:\数据库\KEGG数据库\KEGG通路图")

# (KEGG id, 中文名) —— 英文名以 KEGG 官方为准
CANDIDATES = [
    ("hsa05200", "癌症通路（总览）"),
    ("hsa04216", "铁死亡"),
    ("hsa04142", "溶酶体"),
    ("hsa04141", "内质网蛋白质加工"),
    ("hsa04144", "胞吞作用"),
    ("hsa04146", "过氧化物酶体"),
    ("hsa04330", "Notch信号通路"),
    ("hsa04340", "Hedgehog信号通路"),
    ("hsa04512", "ECM-受体相互作用"),
    ("hsa04612", "抗原加工与呈递"),
    ("hsa04022", "cGMP-PKG信号通路"),
    ("hsa04710", "昼夜节律"),
    ("hsa04714", "产热作用"),
    ("hsa04724", "谷氨酸能突触"),
    ("hsa04727", "GABA能突触"),
    ("hsa04728", "多巴胺能突触"),
    ("hsa04911", "胰岛素分泌"),
    ("hsa04974", "蛋白质消化与吸收"),
    ("hsa04978", "矿物质吸收"),
    ("hsa04979", "胆固醇代谢"),
    ("hsa05012", "帕金森病"),
    ("hsa05014", "肌萎缩侧索硬化"),
    ("hsa05016", "亨廷顿病"),
    ("hsa05410", "肥厚型心肌病"),
    ("hsa05414", "扩张型心肌病"),
    ("hsa04260", "心肌收缩"),
    ("hsa04261", "心肌细胞肾上腺素能信号"),
    ("hsa00100", "类固醇生物合成"),
    ("hsa04975", "脂肪消化与吸收"),
    ("hsa04977", "维生素消化与吸收"),
]


def fname_en(en: str) -> str:
    s = re.sub(r"[^A-Za-z0-9\- ]", "", en).replace(" ", "_")
    return re.sub(r"_+", "_", s)


def main():
    s = requests.Session()
    s.headers.update({"User-Agent": "Mozilla/5.0"})
    official = {}
    r = s.get("https://rest.kegg.jp/list/pathway/hsa", timeout=60)
    r.raise_for_status()
    for line in r.text.splitlines():
        pid, name = line.split("\t", 1)
        official[pid.replace("path:", "")] = name.split(" - Homo sapiens")[0]

    wb = load_workbook(XLSX)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    i_rank = header.index("排名")
    existing = {r[header.index("KEGG编号")] for r in rows[1:]}
    max_rank = max(int(r[i_rank]) for r in rows[1:] if r[i_rank] is not None)

    added = 0
    for pid, cn in CANDIDATES:
        if pid in existing:
            print(f"[skip] {pid} 已在表", flush=True)
            continue
        en = official.get(pid)
        if not en:
            print(f"[warn] {pid} KEGG 无此条目，跳过", flush=True)
            continue
        max_rank += 1
        fen = fname_en(en)
        png = f"{max_rank:03d}_{pid}_{fen}.png"
        gif = f"{max_rank:03d}_{pid}_{cn}_{fen}_动画.gif"
        url = f"https://rest.kegg.jp/get/{pid}/image"
        ws.append([max_rank, pid, cn, en, "待绘制", gif, png, url])
        existing.add(pid)
        added += 1

        dest = OUT / png
        if not dest.exists():
            try:
                rr = s.get(url, timeout=60)
                rr.raise_for_status()
                if rr.content.startswith(b"\x89PNG"):
                    dest.write_bytes(rr.content)
                    print(f"[png] {png} {len(rr.content)//1024}KB", flush=True)
            except Exception as e:
                print(f"[png-fail] {pid} {e}", flush=True)
            time.sleep(0.4)
        print(f"[add] {max_rank} {pid} {cn} / {en}", flush=True)

    if added:
        wb.save(XLSX)
    print(f"done: added={added}, total={max_rank}", flush=True)


if __name__ == "__main__":
    main()
