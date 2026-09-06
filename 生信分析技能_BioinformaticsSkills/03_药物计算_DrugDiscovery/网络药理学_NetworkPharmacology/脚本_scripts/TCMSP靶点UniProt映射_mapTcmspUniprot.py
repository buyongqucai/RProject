# -*- coding: utf-8 -*-
"""TCMSP 靶点名 → UniProt reviewed human 基因简称（第三列）。

从 UniProt REST 拉取 reviewed + organism_id:9606，按蛋白推荐名/别名匹配。
匹配不上的（非人源/过时名）第三列留空，不编造。
"""
from __future__ import annotations

import argparse
import csv
import re
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

UA = {"User-Agent": "TCMSP-UniProt-mapper/1.0 (local research; mailto:local)"}
UNIPROT_SEARCH = "https://rest.uniprot.org/uniprotkb/search"
DEFAULT_TSV = Path(r"D:\数据库\药物数据库\准备文件\_uniprot_reviewed_human.tsv")
PREP = Path(r"D:\数据库\药物数据库\准备文件")

# TCMSP 常用写法 ≠ UniProt 推荐名（仅人源、一对一）
MANUAL = {
    "heat shock protein hsp 90": "HSP90AA1",
    "heat shock protein hsp 90-alpha": "HSP90AA1",
    "dna topoisomerase ii": "TOP2A",
    "coagulation factor xa": "F10",
    "vascular endothelial growth factor a": "VEGFA",
    "beta-secretase": "BACE1",
    "cytochrome p450 51": "CYP51A1",
    "antithrombin-iii precursor": "SERPINC1",
    "mrna of pka catalytic subunit c-alpha": "PRKACA",
    "pka catalytic subunit c-alpha": "PRKACA",
    "phosphatidylinositol-4,5-bisphosphate 3-kinase catalytic subunit, gamma isoform": "PIK3CG",
    "pi3-kinase p110-gamma subunit": "PIK3CG",
    "thrombin": "F2",
    "peroxisome proliferator activated receptor gamma": "PPARG",
    "peroxisome proliferator activated receptor delta": "PPARD",
    "neuronal acetylcholine receptor protein, alpha-7 chain": "CHRNA7",
    "gamma-aminobutyric-acid receptor alpha-2 subunit": "GABRA2",
    "gamma-aminobutyric-acid receptor alpha-3 subunit": "GABRA3",
    "gamma-aminobutyric-acid receptor alpha-5 subunit": "GABRA5",
    "nad(p)h dehydrogenase [quinone] 1": "NQO1",
    "mrna of protein-tyrosine phosphatase, non-receptor type 1": "PTPN1",
    "protein-tyrosine phosphatase, non-receptor type 1": "PTPN1",
    "11-beta-hydroxysteroid dehydrogenase 2": "HSD11B2",
    "dna (cytosine-5)-methyltransferase 3a": "DNMT3A",
    "dna (cytosine-5)-methyltransferase 3b": "DNMT3B",
    "chymotrypsin c": "CTRC",
    "type iv phosphodiesterase": "PDE4A",
    "bifunctional protein ncoat": "MGEA5",
}


def split_uniprot_names(protein_names: str) -> list[str]:
    """按 UniProt 「推荐名 (别名) (别名)」切开，保留推荐名内部括号（如 D(1A)）。"""
    s = (protein_names or "").strip()
    if not s:
        return []
    parts, depth, start = [], 0, 0
    i = 0
    while i < len(s):
        ch = s[i]
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        elif depth == 0 and ch == " " and i + 1 < len(s) and s[i + 1] == "(":
            parts.append(s[start:i].strip())
            start = i + 1
        i += 1
    parts.append(s[start:].strip())
    out, seen = [], set()
    for p in parts:
        p = p.strip()
        if p.startswith("(") and p.endswith(")"):
            p = p[1:-1].strip()
        if p.upper().startswith("EC ") or p.upper().startswith("CLEAVED INTO"):
            break
        key = re.sub(r"\s+", " ", p).lower()
        if p and key not in seen:
            seen.add(key)
            out.append(p)
    return out


def norm(s: str) -> str:
    s = (s or "").strip()
    s = re.sub(r"^mRNA of\s+", "", s, flags=re.I)
    s = s.replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", " ", s)
    return s.lower()


def download_reviewed_human(dest: Path, force: bool = False) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 100_000 and not force:
        return dest
    s = requests.Session()
    s.headers.update(UA)
    params = {
        "query": "(reviewed:true) AND (organism_id:9606)",
        "fields": "accession,id,protein_name,gene_names,gene_primary",
        "format": "tsv",
        "size": "500",
    }
    url, rows, header = UNIPROT_SEARCH, [], None
    page = 0
    while url:
        page += 1
        r = s.get(url, params=params if page == 1 else None, timeout=120)
        r.raise_for_status()
        lines = r.text.splitlines()
        if not lines:
            break
        if header is None:
            header = lines[0]
        rows.extend(lines[1:])
        print(f"  uniprot page {page} +{len(lines)-1} total={len(rows)}", flush=True)
        nxt = None
        link = r.headers.get("Link") or ""
        m = re.search(r'<([^>]+)>;\s*rel="next"', link)
        if m:
            nxt = m.group(1)
        url, params = nxt, None
        time.sleep(0.2)
    dest.write_text(header + "\n" + "\n".join(rows) + "\n", encoding="utf-8")
    print(f"[uniprot] wrote {len(rows)} rows -> {dest}", flush=True)
    return dest


def build_resolver(tsv: Path) -> callable:
    rec_exact: dict[str, set[str]] = {}
    alias_map: dict[str, set[str]] = {}
    with tsv.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        # stream TSV from UniProt uses: Entry, Entry Name, Protein names, Gene Names, Gene Names (primary)
        # search TSV may use same headers
        for row in reader:
            entry = (row.get("Entry") or row.get("accession") or "").strip()
            gene = (row.get("Gene Names (primary)") or row.get("gene_primary")
                    or "").strip()
            if not entry or not gene:
                continue
            pn = row.get("Protein names") or row.get("protein_name") or ""
            aliases = split_uniprot_names(pn)
            if aliases:
                rec_exact.setdefault(norm(aliases[0]), set()).add(gene)
                for a in aliases:
                    alias_map.setdefault(norm(a), set()).add(gene)
            alias_map.setdefault(gene.lower(), set()).add(gene)

    def resolve(target_name: str) -> str:
        nt = norm(target_name)
        if not nt or nt.startswith("unnamed"):
            return ""
        if nt in MANUAL:
            return MANUAL[nt]
        hits = rec_exact.get(nt, set())
        if len(hits) == 1:
            return next(iter(hits))
        hits = alias_map.get(nt, set())
        if len(hits) == 1:
            return next(iter(hits))
        return ""

    return resolve


def fill_target_xlsx(path: Path, resolve) -> tuple[int, int]:
    """写入/覆盖第三列为基因简称。返回 (nrows, n_mapped)。"""
    wb = load_workbook(path)
    ws = wb.active
    n = n_ok = 0
    for row in ws.iter_rows(min_row=1, max_col=2):
        c0, c1 = row[0].value, row[1].value
        if not c0 and not c1:
            continue
        n += 1
        gene = resolve(str(c1 or ""))
        ws.cell(row[0].row, 3, gene)
        if gene:
            n_ok += 1
    wb.save(path)
    return n, n_ok


def map_prep_dir(out_dir: Path, tsv: Path, force_uniprot: bool = False) -> None:
    tsv = download_reviewed_human(tsv, force=force_uniprot)
    resolve = build_resolver(tsv)
    files = sorted(out_dir.glob("TCMSP*有效成分靶点表格.xlsx"))
    tot = ok = 0
    for i, f in enumerate(files, 1):
        n, n_ok = fill_target_xlsx(f, resolve)
        tot += n
        ok += n_ok
        if i % 50 == 0 or i == len(files):
            print(f"[{i}/{len(files)}] mapped {ok}/{tot} pairs", flush=True)
    print(f"[done] files={len(files)} pairs={tot} mapped={ok} empty={tot-ok}", flush=True)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(PREP))
    ap.add_argument("--uniprot-tsv", default=str(DEFAULT_TSV))
    ap.add_argument("--refresh-uniprot", action="store_true")
    args = ap.parse_args()
    map_prep_dir(Path(args.out), Path(args.uniprot_tsv), force_uniprot=args.refresh_uniprot)


if __name__ == "__main__":
    main()
